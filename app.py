from flask import Flask, request, jsonify
from flask_cors import CORS
import PyPDF2
from docx import Document
from pptx import Presentation
import os
from uuid import uuid4

app = Flask(__name__)
CORS(app)  # Google Sites-dan so'rovlarni qabul qilish
UPLOAD_FOLDER = 'uploads'
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# Grok API taqlidi (haqiqiy API bilan almashtiriladi)
def grok_generate_questions(text):
    questions = [
        {
            "type": "multiple_choice",
            "question": "What is the main topic of the document?",
            "options": ["Topic A", "Topic B", "Topic C", "Topic D"],
            "correct": "Topic A"
        },
        {
            "type": "open",
            "question": "Explain the key concept discussed in the document.",
            "answer": ""
        }
    ]
    return questions

# Hujjatdan matn olish
def extract_text_from_file(file_path, file_type):
    text = ""
    if file_type == 'pdf':
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            for page in reader.pages:
                text += page.extract_text() or ""
    elif file_type == 'docx':
        doc = Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
    elif file_type == 'pptx':
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    return text

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    
    file_ext = file.filename.rsplit('.', 1)[1].lower()
    if file_ext not in ['pdf', 'docx', 'pptx']:
        return jsonify({"error": "Unsupported file format"}), 400
    
    file_path = os.path.join(app.config['UPLOAD_FOLDER'], f"{uuid4()}.{file_ext}")
    file.save(file_path)
    
    text = extract_text_from_file(file_path, file_ext)
    questions = grok_generate_questions(text)
    
    os.remove(file_path)
    return jsonify({"questions": questions})

@app.route('/evaluate', methods=['POST'])
def evaluate_answers():
    data = request.json
    questions = data.get('questions', [])
    answers = data.get('answers', [])
    
    results = []
    for q, a in zip(questions, answers):
        if q['type'] == 'multiple_choice':
            score = 1 if a == q['correct'] else 0
        else:
            score = 1 if a.strip() else 0
        results.append({"question": q['question'], "score": score})
    
    return jsonify({"results": results})

if __name__ == '__main__':
    os.makedirs(UPLOAD_FOLDER, exist_ok=True)
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)